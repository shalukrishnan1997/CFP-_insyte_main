#!/usr/bin/env python3
"""Generate docs/letter_templates/letter_template_field_reference.docx for Letter Setup.

Uses the same placeholder catalogue as the admin UI (AVAILABLE_PLACEHOLDERS).

Run from repository root::

    uv run python scripts/build_letter_reference_docx.py
"""

from __future__ import annotations

import os
import sys

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, _REPO_ROOT)
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "responsehandling.settings.test")

import django  # noqa: E402

django.setup()

from letters.view_helpers import AVAILABLE_PLACEHOLDERS  # noqa: E402


def main() -> None:
    from docx import Document
    from docx.shared import Pt

    out_dir = os.path.join(_REPO_ROOT, "docs", "letter_templates")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "letter_template_field_reference.docx")

    doc = Document()
    title = doc.add_heading("Letter template — field reference", level=0)
    for run in title.runs:
        run.font.size = Pt(18)

    intro = doc.add_paragraph()
    intro.add_run(
        "Upload this file only as a reference, or copy sections into your live templates. "
        "Placeholders use docxtpl/Jinja2 syntax. The admin Letter Setup page lists the same "
        "supported variables."
    )

    doc.add_paragraph()
    doc.add_heading("Supported merge fields", level=1)

    category_titles = {
        "donor": "Donor",
        "donation": "Donation",
        "campaign": "Campaign",
        "client": "Client (organisation)",
        "other": "Letter date (when generated)",
    }

    for key, heading in category_titles.items():
        items = AVAILABLE_PLACEHOLDERS.get(key, [])
        if not items:
            continue
        doc.add_heading(heading, level=2)
        for item in items:
            ph = item["placeholder"]
            desc = item.get("description", "")
            p = doc.add_paragraph()
            p.add_run(f"{desc}: ").bold = True
            p.add_run(ph)

    doc.add_page_break()
    doc.add_heading("Condition examples (docxtpl)", level=1)

    doc.add_paragraph(
        "Use docxtpl paragraph-level tags (p-prefix form) so each opening tag, body text, "
        "and closing tag sits in its own Word paragraph. Adjust thresholds and wording "
        "for your charity."
    )

    doc.add_heading("Example A — Gift Aid", level=2)
    doc.add_paragraph("{%p if gift_aid == 'Yes' %}")
    doc.add_paragraph(
        "We have recorded your Gift Aid declaration on this gift. Thank you."
    )
    doc.add_paragraph("{%p endif %}")

    doc.add_heading("Example B — Amount tiers", level=2)
    doc.add_paragraph("{%p if amount_raw >= 100 %}")
    doc.add_paragraph(
        "Thank you for your especially generous support - it makes a real difference."
    )
    doc.add_paragraph("{%p elif amount_raw >= 25 %}")
    doc.add_paragraph("Thank you for your generous gift.")
    doc.add_paragraph("{%p else %}")
    doc.add_paragraph("Thank you for your support.")
    doc.add_paragraph("{%p endif %}")

    doc.add_heading("Example C — Payment completed", level=2)
    doc.add_paragraph("{%p if payment_status == 'completed' %}")
    doc.add_paragraph("Your payment has been received successfully.")
    doc.add_paragraph("{%p endif %}")

    doc.add_heading("Jinja filters (optional)", level=1)
    doc.add_paragraph(
        "Letter generation registers extra filters (see core/letter_tasks.py). "
        "In your template you can pipe values to them using standard Jinja syntax."
    )
    doc.add_paragraph(
        "date_format — use on donation_date_obj; pass a strftime pattern such as "
        "day/month/year with leading zeros."
    )
    doc.add_paragraph(
        "currency — use on amount_raw; optional argument for currency symbol "
        "(default pound)."
    )
    doc.add_paragraph("to_int — whole number from amount_raw.")
    doc.add_paragraph(
        "to_float — decimal string; optional argument for number of places."
    )

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
