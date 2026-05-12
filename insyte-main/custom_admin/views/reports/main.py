"""Report view functions for the admin dashboard.

Views:
    admin_reports: Main reports dashboard with filtering and charts.
    report_export: CSV/Excel export of report data.
"""

import csv
import json
import re
from typing import Any

from django.conf import settings
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.http import HttpRequest, HttpResponse
from django.shortcuts import redirect, render

from audit.utils import log_request_action
from campaigns.models import Campaign
from clients.models import Client
from core.constants import CURRENCY_SYMBOL
from core.utils import restore_session_filters
from responsehandling.permissions import is_authenticated_and_is_staff

from .charts import _prepare_chart_data
from .constants import REPORT_CATEGORIES, REPORT_TYPES, DecimalEncoder
from .generators import (
    _generate_banking_report,
    _generate_campaign_summary_report,
    _generate_credit_card_report,
    _generate_donations_report,
    _generate_gift_aid_report,
    _generate_hgv_report,
    _generate_lgv_report,
    _generate_paying_in_slips_report,
    _generate_payment_method_report,
    _generate_roi_report,
    _generate_unmatched_donors_report,
)
from .helpers import _get_filtered_donations, _parse_date_param
from .scans_zip import report_supports_scan_zip

_REPORT_GENERATORS: dict[str, Any] = {
    "donations": _generate_donations_report,
    "payment_method": _generate_payment_method_report,
    "credit_card": _generate_credit_card_report,
    "banking": _generate_banking_report,
    "gift_aid": _generate_gift_aid_report,
    "roi": _generate_roi_report,
    "hgv": _generate_hgv_report,
    "lgv": _generate_lgv_report,
    "campaign_summary": _generate_campaign_summary_report,
}

# Regex for sanitizing CSV/Excel cell values against formula injection
_FORMULA_PREFIX_RE = re.compile(r"^[=+\-@\t\r]")

_SESSION_KEY = "reports_filters"


def _credit_card_report_headers() -> list[str]:
    """Table/export headers for the credit card report (respects export flag)."""
    headers = [
        "S. No.",
        "Date",
        "Donor Name",
        "URN",
        "Campaign",
        "Amount",
    ]
    if getattr(settings, "ALLOW_CARD_METADATA_EXPORT", True):
        headers.extend(["Card Holder", "Card Number", "Expiry"])
    return headers


_PERSISTENT_PARAMS = (
    "report_type",
    "scope",
    "client",
    "campaign",
    "date_from",
    "date_to",
    "per_page",
)


def _sanitize_cell(value: Any) -> str:
    """Sanitize a cell value to prevent CSV formula injection."""
    s = str(value) if value is not None else ""
    if _FORMULA_PREFIX_RE.match(s):
        return f"'{s}"
    return s


def _save_session_filters(request: HttpRequest) -> None:
    """Persist current GET params to session."""
    filters = {p: request.GET[p] for p in _PERSISTENT_PARAMS if request.GET.get(p)}
    request.session[_SESSION_KEY] = filters


def _build_pagination_query(request: HttpRequest, per_page: int) -> str:
    """Build query string for pagination links."""
    keep = ("report_type", "scope", "client", "campaign", "date_from", "date_to")
    parts = [f"{k}={request.GET[k]}" for k in keep if request.GET.get(k)]
    if per_page != 25:
        parts.append(f"per_page={per_page}")
    return "&".join(parts)


def _campaign_display_name(c: Any) -> str:
    """Return display name for a campaign, with appeal code if present."""
    return f"{c.name} ({c.appeal_code})" if c.appeal_code else c.name


def _build_campaigns_json(campaigns: Any) -> list[dict[str, str]]:
    """Serialize campaigns for frontend JSON."""
    return [
        {
            "id": str(c.id),
            "name": _campaign_display_name(c),
            "client_id": str(c.client_id) if c.client_id else "",
        }
        for c in campaigns
    ]


def _build_report_context(
    request: HttpRequest,
    report_type: str,
    date_from: Any,
    date_to: Any,
    client_id: str,
    campaign_id: str,
    scope: str,
    page: str,
    per_page: int,
) -> dict[str, Any]:
    """Generate paginated report data and chart context."""
    donations = _get_filtered_donations(date_from, date_to, client_id, campaign_id)
    report_data = _run_report_generator(
        report_type, donations, date_from, date_to, client_id
    )
    for i, row in enumerate(report_data):
        row.insert(0, str(i + 1))

    paginated_data, page_num, total_records = _paginate_report(
        report_data, page, per_page
    )
    pagination_query = _build_pagination_query(request, per_page)
    chart_data = _prepare_chart_data(
        report_type,
        donations,
        date_from=date_from,
        date_to=date_to,
        client_id=client_id,
        campaign_id=campaign_id,
        scope=scope,
    )
    cfg = REPORT_TYPES[report_type]
    table_headers = (
        _credit_card_report_headers()
        if report_type == "credit_card"
        else cfg["headers"]
    )
    return {
        "report_title": cfg["title"],
        "report_description": cfg["description"],
        "table_headers": table_headers,
        "report_data": paginated_data.object_list,
        "total_records": total_records,
        "page_start": paginated_data.start_index(),
        "page_end": paginated_data.end_index(),
        "current_page": page_num,
        "has_previous": paginated_data.has_previous(),
        "has_next": paginated_data.has_next(),
        "previous_page": paginated_data.previous_page_number()
        if paginated_data.has_previous()
        else None,
        "next_page": paginated_data.next_page_number()
        if paginated_data.has_next()
        else None,
        "per_page": per_page,
        "pagination_query": pagination_query,
        "chart_data_json": json.dumps(chart_data, cls=DecimalEncoder)
        if chart_data and total_records
        else "null",
        "has_visualization": chart_data is not None and total_records > 0,
        "pdf_row_warning": total_records > 500,
        "scans_zip_supported": report_supports_scan_zip(report_type),
    }


def _run_report_generator(
    report_type: str,
    donations: Any,
    date_from: Any,
    date_to: Any,
    client_id: str,
) -> list:
    """Run the appropriate report generator for a given report type."""
    if report_type == "paying_in_slips":
        return _generate_paying_in_slips_report(
            date_from=date_from,
            date_to=date_to,
            client_id=client_id,
        )
    if report_type == "unmatched_donors":
        return _generate_unmatched_donors_report(
            date_from=date_from,
            date_to=date_to,
            client_id=client_id,
        )
    generator = _REPORT_GENERATORS.get(report_type)
    if not generator:
        return []
    return generator(donations)


def _paginate_report(
    report_data: list, page: str, per_page: int
) -> tuple[Any, int, int]:
    """Paginate report data.

    Returns:
        Tuple of (paginated_data, page_num, total_records).
    """
    total_records = len(report_data)
    paginator = Paginator(report_data, per_page)
    try:
        page_num = int(page)
        paginated_data = paginator.page(page_num)
    except (PageNotAnInteger, ValueError):  # fmt: skip
        page_num = 1
        paginated_data = paginator.page(1)
    except EmptyPage:
        page_num = paginator.num_pages
        paginated_data = paginator.page(paginator.num_pages)
    return paginated_data, page_num, total_records


@is_authenticated_and_is_staff
def admin_reports(request: HttpRequest) -> HttpResponse:
    """Reports dashboard with multiple report types."""
    # Handle filter persistence
    if request.GET.get("clear"):
        request.session.pop(_SESSION_KEY, None)
        return redirect("custom_admin:admin_reports")

    report_type = request.GET.get("report_type", "")

    if not report_type:
        redir = restore_session_filters(
            request, _SESSION_KEY, "custom_admin:admin_reports"
        )
        if redir:
            return redir

    if report_type:
        _save_session_filters(request)

    scope = request.GET.get("scope", "overall")
    client_id = request.GET.get("client", "")
    campaign_id = request.GET.get("campaign", "")
    date_from = _parse_date_param(request.GET.get("date_from"), default_offset_days=30)
    date_to = _parse_date_param(request.GET.get("date_to"))
    page = request.GET.get("page", "1")
    per_page = int(request.GET.get("per_page", "25"))

    all_clients = Client.objects.filter(is_active=True).order_by("name")
    all_campaigns = Campaign.objects.filter(client__is_active=True).order_by("name")

    context: dict[str, Any] = {
        "report_type": report_type,
        "scope": scope,
        "client_id": client_id,
        "campaign_id": campaign_id,
        "date_from": date_from,
        "date_to": date_to,
        "all_clients": all_clients,
        "all_campaigns": all_campaigns,
        "all_clients_json": json.dumps(
            [{"id": str(c.id), "name": c.name} for c in all_clients]
        ),
        "all_campaigns_json": json.dumps(_build_campaigns_json(all_campaigns)),
        "active": "reports",
        "currency_symbol": CURRENCY_SYMBOL,
        "breadcrumbs": [{"name": "Reports", "url": None}],
        "report_count": len(REPORT_TYPES),
        # Category data for the card grid landing page
        "categories": [
            {
                **cat,
                "report_list": [
                    {"key": key, **REPORT_TYPES[key]}
                    for key in cat["reports"]
                    if key in REPORT_TYPES
                ],
            }
            for cat in REPORT_CATEGORIES
        ],
    }

    if report_type and report_type in REPORT_TYPES:
        context.update(
            _build_report_context(
                request,
                report_type,
                date_from,
                date_to,
                client_id,
                campaign_id,
                scope,
                page,
                per_page,
            )
        )

    return render(request, "admin/reports.html", context)


def _auto_fit_columns(ws: Any) -> None:
    """Auto-fit Excel column widths based on cell content."""
    for col_cells in ws.columns:
        max_length = 0
        column_letter = col_cells[0].column_letter
        for cell in col_cells:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except (TypeError, AttributeError):  # fmt: skip
                pass
        ws.column_dimensions[column_letter].width = min(max_length + 2, 50)


def _export_excel(
    report_data: list[list[str]],
    headers: list[str],
    title: str,
    filename: str,
) -> HttpResponse | None:
    """Build an Excel response, or return None if openpyxl is unavailable."""
    try:
        from openpyxl import Workbook  # type: ignore[import-not-found]
        from openpyxl.styles import (  # type: ignore[import-not-found]
            Alignment,
            Font,
            PatternFill,
        )
    except ImportError:
        return None

    wb = Workbook()
    ws = wb.active
    assert ws is not None, "Workbook has no active sheet"
    ws.title = title[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="4472C4", end_color="4472C4", fill_type="solid"
    )
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=_sanitize_cell(header))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row_data in enumerate(report_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            ws.cell(row=row_idx, column=col_idx, value=_sanitize_cell(value))

    _auto_fit_columns(ws)

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response


def _export_csv(
    report_data: list[list[str]],
    headers: list[str],
    filename: str,
) -> HttpResponse:
    """Build a CSV response with cell sanitization."""
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'

    writer = csv.writer(response)
    writer.writerow([_sanitize_cell(h) for h in headers])
    for row in report_data:
        writer.writerow([_sanitize_cell(v) for v in row])
    return response


def _log_export(
    request: HttpRequest,
    title: str,
    export_format: str,
    record_count: int,
) -> None:
    """Write audit log entries for a report export."""
    log_request_action(
        request,
        action="DOWNLOAD",
        model_name="Report",
        object_repr=f"{title} ({export_format.upper()})",
        summary=f"Exported {title} as {export_format.upper()} ({record_count} records)",
    )


@is_authenticated_and_is_staff
def report_export(request: HttpRequest) -> HttpResponse:
    """Export reports to CSV or Excel format."""
    report_type = request.GET.get("report_type", "donations")
    export_format = request.GET.get("format", "csv")
    client_id = request.GET.get("client", "")
    campaign_id = request.GET.get("campaign", "")
    date_from = _parse_date_param(request.GET.get("date_from"), default_offset_days=30)
    date_to = _parse_date_param(request.GET.get("date_to"))

    donations = _get_filtered_donations(date_from, date_to, client_id, campaign_id)
    report_data = _run_report_generator(
        report_type, donations, date_from, date_to, client_id
    )
    for i, row in enumerate(report_data):
        row.insert(0, str(i + 1))

    config = REPORT_TYPES.get(report_type, REPORT_TYPES["donations"])
    headers: list[str] = (
        _credit_card_report_headers()
        if report_type == "credit_card"
        else config["headers"]
    )
    filename = f"{report_type}_report_{date_from}_to_{date_to}"

    if export_format == "excel":
        excel_resp = _export_excel(report_data, headers, config["title"], filename)
        if excel_resp is not None:
            _log_export(request, config["title"], export_format, len(report_data))
            return excel_resp

    response = _export_csv(report_data, headers, filename)
    _log_export(request, config["title"], export_format, len(report_data))
    return response
