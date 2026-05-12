"""Unit tests for letter batch task helpers."""

from pathlib import Path
from typing import Any
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
from docx import Document

from campaigns.models import CampaignDataFile
from donations.models import Donation
from donors.models import DataFileDonor
from letters.models import LetterBatch, LetterTemplate
from letters.tasks import (
    _build_generation_groups,
    _process_letter_chunk,
    _run_letter_generation,
    build_letter_generation_queryset,
    generate_merged_document,
    reset_failed_donations,
)
from tests.factories import (
    CampaignFactory,
    DonationBatchFactory,
    DonationFactory,
    DonorFactory,
    UserFactory,
)


@pytest.mark.django_db()
class TestLetterTaskHelpers:
    """Tests for core letter task helper behavior."""

    def test_build_generation_groups_routes_approved_to_thanks(self) -> None:
        """Approved donations are routed to the thank-you template."""
        campaign = CampaignFactory()
        approved = DonationFactory(
            campaign=campaign, qa_status=Donation.QA_STATUS_APPROVED
        )
        DonationFactory(campaign=campaign, qa_status=Donation.QA_STATUS_REJECTED)
        donations_qs = Donation.objects.filter(campaign=campaign)

        groups = _build_generation_groups(
            donations_qs=donations_qs,
            template_path="thanks.docx",
            failure_template_path=None,
        )

        assert len(groups) == 1
        ids, template_path, label = groups[0]
        assert template_path == "thanks.docx"
        assert label == "THANKS"
        assert ids == [approved.id]

    def test_build_generation_groups_skips_rejected_without_issue_template(
        self,
    ) -> None:
        """Rejected donations are skipped when no issue template is configured."""
        campaign = CampaignFactory()
        DonationFactory(campaign=campaign, qa_status=Donation.QA_STATUS_REJECTED)
        donations_qs = Donation.objects.filter(campaign=campaign)

        groups = _build_generation_groups(
            donations_qs=donations_qs,
            template_path="thanks.docx",
            failure_template_path=None,
        )

        assert groups == []

    def test_build_generation_groups_routes_both_when_templates_present(self) -> None:
        """Approved donations go to thanks and rejected donations go to issue."""
        campaign = CampaignFactory()
        approved = DonationFactory(
            campaign=campaign, qa_status=Donation.QA_STATUS_APPROVED
        )
        rejected = DonationFactory(
            campaign=campaign, qa_status=Donation.QA_STATUS_REJECTED
        )
        donations_qs = Donation.objects.filter(campaign=campaign)

        groups = _build_generation_groups(
            donations_qs=donations_qs,
            template_path="thanks.docx",
            failure_template_path="issue.docx",
        )

        assert len(groups) == 2
        thanks_ids, thanks_path, thanks_label = groups[0]
        issue_ids, issue_path, issue_label = groups[1]
        assert thanks_label == "THANKS"
        assert thanks_path == "thanks.docx"
        assert thanks_ids == [approved.id]
        assert issue_label == "ISSUE"
        assert issue_path == "issue.docx"
        assert issue_ids == [rejected.id]

    def test_build_letter_generation_queryset_excludes_non_financial_donor_update(
        self,
    ) -> None:
        """Phone-intake donor-update rows must never queue a thank-you letter.

        ``process_phone_non_financial_intake`` stamps ``letter_voided_at``
        and ``letter_status='skipped'``; both belt and suspenders should
        keep the row out of the eligible queryset.
        """
        from django.utils import timezone

        campaign = CampaignFactory()
        donor_update = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            payment_method="non_financial",
            amount=0,
            letter_status="skipped",
            letter_void_reason="Non-financial donor update",
            letter_voided_at=timezone.now(),
        )
        thanks_eligible = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )

        donations_qs = build_letter_generation_queryset(
            campaign=campaign,
            donation_filter="all",
            regenerate_mode=False,
        )

        assert donor_update not in donations_qs
        assert thanks_eligible in donations_qs

    def test_build_letter_generation_queryset_excludes_pending_and_flagged(
        self,
    ) -> None:
        """Donations in qa_status pending or flagged are not eligible for letters."""
        campaign = CampaignFactory()
        approved = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )
        rejected = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_REJECTED,
            letter_status="pending",
        )
        pending = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_PENDING,
            letter_status="pending",
        )
        flagged = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_FLAGGED,
            letter_status="pending",
        )

        donations_qs = build_letter_generation_queryset(
            campaign=campaign,
            donation_filter="all",
            regenerate_mode=False,
        )

        assert approved in donations_qs
        assert rejected in donations_qs
        assert pending not in donations_qs
        assert flagged not in donations_qs

    def test_build_letter_generation_queryset_excludes_deceased_house_donor(
        self,
    ) -> None:
        """Postal suppression excludes deceased house-file donors from letters."""
        campaign = CampaignFactory()
        deceased_donor = DonorFactory(contact_status="deceased")
        active_donor = DonorFactory()
        excluded = DonationFactory(
            campaign=campaign,
            donor=deceased_donor,
            qa_status=Donation.QA_STATUS_APPROVED,
        )
        included = DonationFactory(
            campaign=campaign,
            donor=active_donor,
            qa_status=Donation.QA_STATUS_APPROVED,
        )

        donations_qs = build_letter_generation_queryset(
            campaign=campaign,
            donation_filter="all",
            regenerate_mode=False,
        )

        assert excluded not in donations_qs
        assert included in donations_qs

    def test_build_letter_generation_queryset_excludes_gone_away_data_file_donor(
        self,
    ) -> None:
        """Postal suppression excludes gone-away data-file donors from letters."""
        campaign = CampaignFactory()
        data_file = CampaignDataFile.objects.create(
            campaign=campaign,
            created_by=campaign.created_by,
        )
        blocked = DataFileDonor.objects.create(
            data_file=data_file,
            client=campaign.client,
            urn="DFD0001",
            first_name="Blocked",
            last_name="Supporter",
            contact_status=DataFileDonor.CONTACT_STATUS_GONE_AWAY,
            created_by=campaign.created_by,
        )
        eligible = DataFileDonor.objects.create(
            data_file=data_file,
            client=campaign.client,
            urn="DFD0002",
            first_name="Eligible",
            last_name="Supporter",
            created_by=campaign.created_by,
        )
        excluded = DonationFactory(
            campaign=campaign,
            donor=None,
            donor_source="data_file",
            data_file_donor=blocked,
            qa_status=Donation.QA_STATUS_APPROVED,
        )
        included = DonationFactory(
            campaign=campaign,
            donor=None,
            donor_source="data_file",
            data_file_donor=eligible,
            qa_status=Donation.QA_STATUS_APPROVED,
        )

        donations_qs = build_letter_generation_queryset(
            campaign=campaign,
            donation_filter="all",
            regenerate_mode=False,
        )

        assert excluded not in donations_qs
        assert included in donations_qs

    def test_build_letter_generation_queryset_scopes_to_source_batch(self) -> None:
        """Scoped runs include both approved and rejected donations from a source batch."""
        campaign = CampaignFactory()
        source_batch = DonationBatchFactory(campaign=campaign)
        other_batch = DonationBatchFactory(campaign=campaign)

        approved_in_source = DonationFactory(
            campaign=campaign,
            batch=source_batch,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )
        rejected_in_source = DonationFactory(
            campaign=campaign,
            batch=source_batch,
            qa_status=Donation.QA_STATUS_REJECTED,
            letter_status="pending",
        )
        approved_in_other_batch = DonationFactory(
            campaign=campaign,
            batch=other_batch,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )

        donations_qs = build_letter_generation_queryset(
            campaign=campaign,
            donation_filter="all",
            regenerate_mode=False,
            source_donation_batch_id=source_batch.id,
        )

        assert approved_in_source in donations_qs
        assert rejected_in_source in donations_qs
        assert approved_in_other_batch not in donations_qs

    def test_run_letter_generation_completes_when_no_eligible_groups(
        self,
        settings: Any,
        tmp_path: Path,
    ) -> None:
        """Batch completes with no files if only rejected donations exist and no issue template is configured."""
        settings.STORAGES = {
            "default": {
                "BACKEND": "django.core.files.storage.FileSystemStorage",
                "OPTIONS": {"location": str(tmp_path)},
            }
        }
        settings.MEDIA_ROOT = str(tmp_path)

        user = UserFactory()
        campaign = CampaignFactory(created_by=user)
        DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_REJECTED,
            letter_status="pending",
        )

        template = LetterTemplate.objects.create(
            campaign=campaign,
            name="Thanks Template",
            template_type="thank_you",
            file="uploads/letter_templates/thanks.docx",
            created_by=user,
        )
        batch = LetterBatch.objects.create(
            campaign=campaign,
            template=template,
            batch_number=1,
            status="pending",
            created_by=user,
        )

        result = _run_letter_generation(
            batch=batch,
            task=Mock(),
            batch_id=str(batch.id),
            donations_qs=build_letter_generation_queryset(
                campaign=campaign,
                donation_filter="all",
                regenerate_mode=False,
            ),
        )

        batch.refresh_from_db()
        assert result["success"] is True
        assert result["generated_count"] == 0
        assert result["failed_count"] == 0
        assert batch.status == "completed"
        assert batch.total_letters == 0

    def test_process_letter_chunk_updates_exact_success_and_failure_statuses(
        self,
        monkeypatch: pytest.MonkeyPatch,
        tmp_path: Path,
    ) -> None:
        """Only successful IDs become generated; remaining IDs are marked failed."""
        user = UserFactory()
        campaign = CampaignFactory(created_by=user)
        first = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )
        second = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )

        template = LetterTemplate.objects.create(
            campaign=campaign,
            name="Thanks Template",
            template_type="thank_you",
            file="uploads/letter_templates/thanks.docx",
            created_by=user,
        )
        batch = LetterBatch.objects.create(
            campaign=campaign,
            template=template,
            batch_number=5,
            status="processing",
            created_by=user,
        )

        def fake_generate_merged_document(
            template_path: str,
            donations: list[Any],
            output_path: str,
        ) -> tuple[int, int, list[str], list[Any]]:
            return 1, 1, ["failed second donation"], [first.id]

        monkeypatch.setattr(
            "letters.tasks.generate_merged_document",
            fake_generate_merged_document,
        )

        generated, failed, errors, output_path = _process_letter_chunk(
            chunk_ids=[first.id, second.id],
            group_template_path="unused.docx",
            output_dir=tmp_path,
            batch=batch,
            group_label="THANKS",
            file_number=1,
        )

        first.refresh_from_db()
        second.refresh_from_db()

        assert generated == 1
        assert failed == 1
        assert errors == ["failed second donation"]
        assert output_path is not None
        assert first.letter_status == "generated"
        assert first.letter_batch_id == batch.id
        assert first.letter_generated_at is not None
        assert second.letter_status == "failed"
        assert second.letter_batch_id == batch.id

    def test_generate_merged_document_keeps_single_section_with_page_breaks(
        self,
        tmp_path: Path,
    ) -> None:
        """Merged output keeps one final section and one page break per extra letter."""
        pytest.importorskip("docxtpl")

        template_path = tmp_path / "template.docx"
        output_path = tmp_path / "merged.docx"

        template_doc = Document()
        template_doc.add_paragraph("Dear {{ donor_first_name }},")
        template_doc.add_paragraph(
            "Thank you for your donation of {{ amount_formatted }}."
        )
        template_doc.save(str(template_path))

        campaign = CampaignFactory()
        first_donor = DonorFactory(first_name="Alice", last_name="Jones")
        second_donor = DonorFactory(first_name="Bob", last_name="Smith")
        first = DonationFactory(
            campaign=campaign,
            donor=first_donor,
            qa_status=Donation.QA_STATUS_APPROVED,
        )
        second = DonationFactory(
            campaign=campaign,
            donor=second_donor,
            qa_status=Donation.QA_STATUS_APPROVED,
        )

        generated, failed, errors, successful_ids = generate_merged_document(
            template_path=str(template_path),
            donations=[first, second],
            output_path=str(output_path),
        )

        assert generated == 2
        assert failed == 0
        assert errors == []
        assert successful_ids == [first.id, second.id]

        with ZipFile(output_path) as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")

        assert document_xml.count('w:type="page"') == 1
        assert document_xml.count("<w:sectPr") == 1
        assert "Alice" in document_xml
        assert "Bob" in document_xml


# ═══════════════════════════════════════════════════════════════
# Banking-reversal cascade guards on letter querysets
# ═══════════════════════════════════════════════════════════════


@pytest.mark.django_db()
class TestBuildLetterGenerationQuerysetCascadeGuards:
    """Ensure build_letter_generation_queryset excludes donations whose
    underlying payment has been invalidated by the banking cascade or
    Stripe-side terminal states. Without this, a regenerate run could
    silently re-issue a letter for a donation that was reversed after the
    earlier letter went out.
    """

    def test_excludes_reversed_donations(self) -> None:
        """payment_status='reversed' is excluded from eligible donations."""
        campaign = CampaignFactory()
        completed = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
            payment_status=Donation.PAYMENT_STATUS_COMPLETED,
        )
        reversed_donation = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
            payment_status=Donation.PAYMENT_STATUS_REVERSED,
        )

        eligible = build_letter_generation_queryset(
            campaign=campaign, donation_filter="all", regenerate_mode=False
        )

        assert completed in eligible
        assert reversed_donation not in eligible

    def test_excludes_voided_donations_in_regenerate_mode(self) -> None:
        """letter_voided_at IS NOT NULL is excluded even when regenerating.

        Regenerate-mode normally drops the ``letter_status='pending'``
        gate so already-generated letters can be re-emitted with a fresh
        template. The cascade-void exclusion has to apply in regenerate
        mode too — otherwise a re-run after a banking reversal would
        silently re-issue an invalidated letter.
        """
        from django.utils import timezone

        campaign = CampaignFactory()
        regenerable = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="generated",
            payment_status=Donation.PAYMENT_STATUS_COMPLETED,
        )
        voided = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="sent",
            payment_status=Donation.PAYMENT_STATUS_REVERSED,
            letter_voided_at=timezone.now(),
            letter_void_reason="Cheque bounced",
        )

        eligible = build_letter_generation_queryset(
            campaign=campaign, donation_filter="all", regenerate_mode=True
        )

        assert regenerable in eligible
        assert voided not in eligible


@pytest.mark.django_db()
class TestResetFailedDonationsCascadeGuards:
    """Ensure reset_failed_donations skips donations the cascade has
    invalidated, but still resets legitimate retry candidates including
    donations whose underlying card payment is pending capture.
    """

    def test_skips_reversed(self) -> None:
        """A failed letter on a reversed donation stays failed, not reset."""
        campaign = CampaignFactory()
        donation = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="failed",
            payment_status=Donation.PAYMENT_STATUS_REVERSED,
        )

        result = reset_failed_donations(str(campaign.id))

        assert result["reset_count"] == 0
        donation.refresh_from_db()
        assert donation.letter_status == "failed"

    def test_skips_voided(self) -> None:
        """A failed letter with letter_voided_at set is left alone."""
        from django.utils import timezone

        campaign = CampaignFactory()
        donation = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="failed",
            payment_status=Donation.PAYMENT_STATUS_COMPLETED,
            letter_voided_at=timezone.now(),
            letter_void_reason="Cheque bounced",
        )

        result = reset_failed_donations(str(campaign.id))

        assert result["reset_count"] == 0
        donation.refresh_from_db()
        assert donation.letter_status == "failed"

    def test_keeps_pending_payment_resettable(self) -> None:
        """A failed letter on a pending-payment donation is still reset.

        Regression guard for the advisor's "don't narrow to completed"
        feedback: card donations awaiting capture have
        ``payment_status='pending'`` but their failed letters are
        legitimate retry candidates once the capture lands.
        """
        campaign = CampaignFactory()
        donation = DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="failed",
            payment_status=Donation.PAYMENT_STATUS_PENDING,
        )

        result = reset_failed_donations(str(campaign.id))

        assert result["reset_count"] == 1
        donation.refresh_from_db()
        assert donation.letter_status == "pending"
        assert donation.letter_generated_at is None
        assert donation.letter_batch is None
