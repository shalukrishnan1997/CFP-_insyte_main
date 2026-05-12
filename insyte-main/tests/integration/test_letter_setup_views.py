"""Integration tests for the letters admin views."""

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client as DjangoClient
from django.urls import reverse
from docx import Document

from donations.models import Donation
from letters.models import LetterBatch, LetterTemplate
from tests.factories import CampaignFactory, DonationFactory


def _create_letter_template(
    campaign: Any,
    created_by: Any,
    *,
    name: str,
    template_type: str,
    is_active: bool = True,
) -> LetterTemplate:
    """Create a campaign-scoped letter template for view tests."""
    return LetterTemplate.objects.create(
        campaign=campaign,
        name=name,
        template_type=template_type,
        file=f"uploads/letter_templates/{template_type}.docx",
        is_active=is_active,
        created_by=created_by,
    )


def _build_docx_upload(*paragraphs: str) -> SimpleUploadedFile:
    """Build an in-memory DOCX file for upload tests."""
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    buffer = BytesIO()
    document.save(buffer)
    return SimpleUploadedFile(
        "template.docx",
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )


@pytest.mark.django_db()
class TestLetterPrintConsole:
    """Tests for the print operator's single-page console."""

    def test_requires_client_and_campaign_selection(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Landing on the console with nothing selected prompts for a client."""
        response = authenticated_client.get(
            reverse("custom_admin:letter_print_console")
        )
        assert response.status_code == 200
        assert b"Pick a client to begin." in response.content

    def test_shows_ready_counts_once_campaign_selected(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Selecting a campaign exposes approved/rejected pending counts."""
        campaign = CampaignFactory(created_by=staff_user)
        _create_letter_template(
            campaign, staff_user, name="Thanks", template_type="thank_you"
        )
        DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )
        DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_REJECTED,
            letter_status="pending",
        )

        response = authenticated_client.get(
            reverse("custom_admin:letter_print_console"),
            {"client": str(campaign.client.id), "campaign": str(campaign.id)},
        )

        assert response.status_code == 200
        assert response.context["approved_pending"] == 1
        assert response.context["rejected_pending"] == 1
        assert response.context["active_thanks_template"] is not None


@pytest.mark.django_db()
class TestLetterSetupCampaign:
    """Tests for the per-campaign template management page."""

    def test_groups_templates_by_type(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Campaign page groups thank-you and issue templates separately."""
        campaign = CampaignFactory(created_by=staff_user)
        thank_template = _create_letter_template(
            campaign, staff_user, name="Thanks", template_type="thank_you"
        )
        issue_template = _create_letter_template(
            campaign, staff_user, name="Issue", template_type="issue"
        )

        response = authenticated_client.get(
            reverse(
                "custom_admin:letter_setup_campaign",
                kwargs={"campaign_id": campaign.id},
            )
        )

        assert response.status_code == 200
        assert list(response.context["thanks_templates"]) == [thank_template]
        assert list(response.context["issue_templates"]) == [issue_template]
        assert response.context["active_thanks_template"] == thank_template
        assert response.context["active_issue_template"] == issue_template

    def test_authoring_card_renders_literal_docxtpl_syntax(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """The syntax-guide card must render literal `{{...}}` / `{% ... %}` to users."""
        campaign = CampaignFactory(created_by=staff_user)
        _create_letter_template(
            campaign, staff_user, name="Thanks", template_type="thank_you"
        )

        response = authenticated_client.get(
            reverse(
                "custom_admin:letter_setup_campaign",
                kwargs={"campaign_id": campaign.id},
            )
        )
        html = response.content.decode()

        assert response.status_code == 200
        # Django {% templatetag %} must produce literal braces, not be evaluated.
        assert "{{ donor_full_name }}" in html
        assert '{% if gift_aid == "Yes" %}' in html
        assert "{% endif %}" in html
        assert "{% elif amount_raw >= 100 %}" in html
        assert 'donation_date_obj|date_format("%d %B %Y")' in html
        # Reference-guide download + Preview button must be wired up.
        assert reverse("custom_admin:letter_reference_guide") in html
        assert (
            reverse(
                "custom_admin:preview_active_template",
                kwargs={"campaign_id": campaign.id, "template_type": "thank_you"},
            )
            in html
        )
        # Dropzone hint advertises the 10 MB limit.
        assert "Max 10 MB" in html

    def test_add_letter_template_get_redirects_to_campaign_workspace(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """GET on the upload URL redirects into the single-page workspace."""
        campaign = CampaignFactory(created_by=staff_user)

        response = authenticated_client.get(
            reverse(
                "custom_admin:add_letter_template",
                kwargs={"campaign_id": campaign.id},
            )
        )

        assert response.status_code == 302
        assert response.url.endswith(
            f"/letter-setup/campaign/{campaign.id}/?section=upload"
        )

    def test_add_letter_template_rejects_invalid_upload(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Non-DOCX uploads are rejected by the upload form."""
        campaign = CampaignFactory(created_by=staff_user)
        invalid_upload = SimpleUploadedFile("template.txt", b"plain text")

        response = authenticated_client.post(
            reverse(
                "custom_admin:add_letter_template",
                kwargs={"campaign_id": campaign.id},
            ),
            {
                "template_name": "Invalid Template",
                "template_type": "thank_you",
                "file": invalid_upload,
            },
        )

        assert response.status_code == 200
        assert LetterTemplate.objects.filter(campaign=campaign).count() == 0
        assert b"Only .docx files are supported." in response.content

    def test_add_letter_template_deactivates_previous_of_same_type(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Uploading a new template deactivates the previous active one of the same type."""
        campaign = CampaignFactory(created_by=staff_user)
        previous = _create_letter_template(
            campaign, staff_user, name="Previous Thanks", template_type="thank_you"
        )

        new_upload = _build_docx_upload("Dear {{ donor_first_name }}")
        response = authenticated_client.post(
            reverse(
                "custom_admin:add_letter_template",
                kwargs={"campaign_id": campaign.id},
            ),
            {
                "template_name": "New Thanks",
                "template_type": "thank_you",
                "file": new_upload,
            },
        )

        assert response.status_code == 302
        previous.refresh_from_db()
        assert previous.is_active is False

        new_template = LetterTemplate.objects.get(campaign=campaign, name="New Thanks")
        assert new_template.is_active is True

    def test_add_letter_template_renders_upload_section_on_validation_failure(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Template inspection failures keep the user on the upload workspace section."""
        campaign = CampaignFactory(created_by=staff_user)
        docx_upload = _build_docx_upload("Dear {{ donor_first_name }}")

        monkeypatch.setattr(
            "letters.admin_views.inspect_docx_template",
            lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("Broken tags")),
        )

        response = authenticated_client.post(
            reverse(
                "custom_admin:add_letter_template",
                kwargs={"campaign_id": campaign.id},
            ),
            {
                "template_name": "Broken Template",
                "template_type": "thank_you",
                "file": docx_upload,
            },
        )

        assert response.status_code == 200
        assert LetterTemplate.objects.filter(campaign=campaign).count() == 0
        assert b"Template validation failed." in response.content
        assert response.context["focus_section"] == "upload"


@pytest.mark.django_db()
class TestGenerateLetter:
    """Tests for the generate_letter POST handler."""

    def test_uses_active_templates_from_campaign(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Generate automatically uses the campaign's active thanks/issue templates."""
        campaign = CampaignFactory(created_by=staff_user)
        thanks = _create_letter_template(
            campaign, staff_user, name="Thanks", template_type="thank_you"
        )
        issue = _create_letter_template(
            campaign, staff_user, name="Issue", template_type="issue"
        )

        DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_APPROVED,
            letter_status="pending",
        )
        DonationFactory(
            campaign=campaign,
            qa_status=Donation.QA_STATUS_REJECTED,
            letter_status="pending",
        )

        monkeypatch.setattr(
            "letters.admin_views.generate_letter_batch_task.delay",
            lambda batch_id: SimpleNamespace(id=f"task-{batch_id}"),
        )

        response = authenticated_client.post(
            reverse(
                "custom_admin:generate_letter",
                kwargs={"campaign_id": campaign.id},
            ),
            {
                "letters_per_file": "100",
                "donation_filter": "all",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        assert response.status_code == 200
        payload = response.json()
        batch = LetterBatch.objects.get(id=payload["batch_id"])

        assert payload["total_letters"] == 2
        assert batch.template == thanks
        assert batch.failure_template == issue

    def test_uses_filtered_total_for_hgv(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """Batch totals reflect the HGV filter and regenerate mode."""
        campaign = CampaignFactory(
            created_by=staff_user,
            hgv_amount=Decimal("100.00"),
        )
        _create_letter_template(
            campaign, staff_user, name="Thanks", template_type="thank_you"
        )

        DonationFactory(
            campaign=campaign,
            amount=Decimal("50.00"),
            letter_status="pending",
            qa_status=Donation.QA_STATUS_APPROVED,
        )
        DonationFactory(
            campaign=campaign,
            amount=Decimal("250.00"),
            letter_status="pending",
            qa_status=Donation.QA_STATUS_APPROVED,
        )
        DonationFactory(
            campaign=campaign,
            amount=Decimal("300.00"),
            letter_status="generated",
            qa_status=Donation.QA_STATUS_APPROVED,
        )

        monkeypatch.setattr(
            "letters.admin_views.generate_letter_batch_task.delay",
            lambda batch_id: SimpleNamespace(id=f"task-{batch_id}"),
        )

        response = authenticated_client.post(
            reverse(
                "custom_admin:generate_letter",
                kwargs={"campaign_id": campaign.id},
            ),
            {
                "letters_per_file": "100",
                "donation_filter": "only_hgv",
                "regenerate_mode": "on",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        assert response.status_code == 200
        payload = response.json()
        batch = LetterBatch.objects.get(id=payload["batch_id"])

        assert payload["total_letters"] == 2
        assert batch.total_letters == 2
        assert batch.regenerate_mode is True
        assert batch.donation_filter == "only_hgv"

    def test_returns_error_when_no_thanks_template(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Generation fails fast with a clear message if no thanks template is configured."""
        campaign = CampaignFactory(created_by=staff_user)

        response = authenticated_client.post(
            reverse(
                "custom_admin:generate_letter",
                kwargs={"campaign_id": campaign.id},
            ),
            {
                "letters_per_file": "100",
                "donation_filter": "all",
            },
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )

        assert response.status_code == 400
        body = response.json()
        assert body["success"] is False
        assert "thank-you template" in body["error"]

    def test_invalid_letters_per_file_rerenders_workspace(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Invalid numeric input keeps the generation form values in the workspace."""
        campaign = CampaignFactory(created_by=staff_user)
        _create_letter_template(
            campaign, staff_user, name="Thanks", template_type="thank_you"
        )

        response = authenticated_client.post(
            reverse(
                "custom_admin:generate_letter",
                kwargs={"campaign_id": campaign.id},
            ),
            {
                "letters_per_file": "abc",
                "donation_filter": "only_hgv",
            },
        )

        assert response.status_code == 200
        assert b"Letters per file must be a whole number." in response.content
        assert response.context["focus_section"] == "generate"
        assert response.context["generation_values"]["donation_filter"] == "only_hgv"


def _create_active_template_with_bytes(
    campaign: Any,
    created_by: Any,
    *,
    template_type: str,
    body: str,
) -> LetterTemplate:
    """Create an active LetterTemplate backed by real DOCX bytes on disk."""
    return LetterTemplate.objects.create(
        campaign=campaign,
        name=f"Active {template_type}",
        template_type=template_type,
        file=_build_docx_upload(body),
        is_active=True,
        created_by=created_by,
    )


@pytest.mark.django_db()
class TestServeReferenceGuide:
    """Tests for the docxtpl reference-guide download view."""

    def test_returns_reference_docx_with_attachment_disposition(
        self,
        authenticated_client: DjangoClient,
    ) -> None:
        """Staff user gets a DOCX attachment with the canonical filename."""
        response = authenticated_client.get(
            reverse("custom_admin:letter_reference_guide")
        )

        assert response.status_code == 200
        assert response["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert (
            'attachment; filename="letter_template_field_reference.docx"'
            in response["Content-Disposition"]
        )
        # Streaming body is non-empty
        body = b"".join(response.streaming_content)
        assert len(body) > 0

    def test_404_when_reference_file_missing(
        self,
        authenticated_client: DjangoClient,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """If the reference DOCX isn't on disk, the view 404s cleanly."""
        monkeypatch.setattr(
            "letters.admin_views.REFERENCE_GUIDE_PATH",
            "/tmp/does-not-exist-ref.docx",
        )

        response = authenticated_client.get(
            reverse("custom_admin:letter_reference_guide")
        )

        assert response.status_code == 404

    def test_redirects_unauthenticated(self, client: DjangoClient) -> None:
        """Anonymous users are sent to login, not served the DOCX."""
        response = client.get(reverse("custom_admin:letter_reference_guide"))

        assert response.status_code == 302
        assert "/login" in response.url


@pytest.mark.django_db()
class TestPreviewActiveTemplate:
    """Tests for the sample-render preview view."""

    def test_renders_active_thank_you_template_to_docx(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Preview returns a rendered DOCX for the active thank-you template."""
        campaign = CampaignFactory(created_by=staff_user)
        _create_active_template_with_bytes(
            campaign,
            staff_user,
            template_type="thank_you",
            body="Dear {{ donor_full_name }}, thanks for {{ amount_formatted }}.",
        )

        response = authenticated_client.get(
            reverse(
                "custom_admin:preview_active_template",
                kwargs={"campaign_id": campaign.id, "template_type": "thank_you"},
            )
        )

        assert response.status_code == 200
        assert response["Content-Type"] == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert "_thank_you_sample.docx" in response["Content-Disposition"]
        body = b"".join(response.streaming_content)
        assert len(body) > 0
        # The rendered DOCX should contain the substituted donor name from the
        # sample validation context (Alex Taylor), not the raw Jinja placeholder.
        rendered = Document(BytesIO(body))
        rendered_text = "\n".join(p.text for p in rendered.paragraphs)
        assert "Alex Taylor" in rendered_text
        assert "{{" not in rendered_text

    def test_404_when_template_type_unknown(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Bogus template_type values are rejected with 404."""
        campaign = CampaignFactory(created_by=staff_user)

        response = authenticated_client.get(
            reverse(
                "custom_admin:preview_active_template",
                kwargs={"campaign_id": campaign.id, "template_type": "bogus"},
            )
        )

        assert response.status_code == 404

    def test_redirects_with_error_when_no_active_template(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
    ) -> None:
        """Campaign with no active template of the type redirects with an error message."""
        campaign = CampaignFactory(created_by=staff_user)

        response = authenticated_client.get(
            reverse(
                "custom_admin:preview_active_template",
                kwargs={"campaign_id": campaign.id, "template_type": "issue"},
            ),
            follow=True,
        )

        assert response.status_code == 200
        messages_list = [str(m) for m in response.context["messages"]]
        assert any("No active issue template" in m for m in messages_list)

    def test_redirects_with_error_when_render_fails(
        self,
        authenticated_client: DjangoClient,
        staff_user: Any,
        monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """A render exception surfaces as a messages.error, not a 500."""
        campaign = CampaignFactory(created_by=staff_user)
        _create_active_template_with_bytes(
            campaign,
            staff_user,
            template_type="thank_you",
            body="Dear {{ donor_full_name }}",
        )

        monkeypatch.setattr(
            "letters.admin_views.DocxTemplate",
            lambda *_args, **_kwargs: (_ for _ in ()).throw(
                ValueError("Simulated render failure")
            ),
        )

        response = authenticated_client.get(
            reverse(
                "custom_admin:preview_active_template",
                kwargs={"campaign_id": campaign.id, "template_type": "thank_you"},
            ),
            follow=True,
        )

        assert response.status_code == 200
        messages_list = [str(m) for m in response.context["messages"]]
        assert any("Preview render failed" in m for m in messages_list)
